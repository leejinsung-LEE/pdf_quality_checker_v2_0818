# -*- coding: utf-8 -*-
"""
데이터 내보내기 및 유지보수 처리
"""

import json
import csv
from pathlib import Path
from datetime import datetime, timedelta
from typing import Optional, List
import logging

from .models import ProcessHistory
from .database import DatabaseManager
from .search_engine import SearchEngine


class ExportHandler:
    """내보내기 및 유지보수 처리기"""
    
    def __init__(self, db_manager: DatabaseManager, 
                 search_engine: SearchEngine,
                 logger: Optional[logging.Logger] = None):
        """
        초기화
        
        Args:
            db_manager: 데이터베이스 관리자
            search_engine: 검색 엔진
            logger: 로거
        """
        self.db = db_manager
        self.search = search_engine
        self.logger = logger or logging.getLogger(__name__)
    
    def export_history(self, output_path: Path, 
                       format: str = "csv",
                       start_date: Optional[datetime] = None,
                       end_date: Optional[datetime] = None,
                       **search_filters) -> bool:
        """
        이력 내보내기
        
        Args:
            output_path: 출력 파일 경로
            format: 출력 형식 (csv, json, excel)
            start_date: 시작 날짜
            end_date: 종료 날짜
            **search_filters: 추가 검색 필터
            
        Returns:
            bool: 성공 여부
        """
        # 이력 검색
        histories = self.search.search_history(
            start_date=start_date,
            end_date=end_date,
            **search_filters
        )
        
        if not histories:
            self.logger.warning("내보낼 이력이 없습니다")
            return False
        
        try:
            if format == "csv":
                return self._export_to_csv(histories, output_path)
            elif format == "json":
                return self._export_to_json(histories, output_path)
            elif format == "excel":
                return self._export_to_excel(histories, output_path)
            else:
                self.logger.error(f"지원하지 않는 형식: {format}")
                return False
                
        except Exception as e:
            self.logger.error(f"이력 내보내기 실패: {e}")
            return False
    
    def _export_to_csv(self, histories: List[ProcessHistory], output_path: Path) -> bool:
        """CSV 형식으로 내보내기"""
        try:
            with open(output_path, 'w', newline='', encoding='utf-8-sig') as f:
                if histories:
                    # 헤더 작성
                    fieldnames = [
                        'id', 'file_name', 'file_path', 'file_size',
                        'process_type', 'status', 'quality_score',
                        'issues_found', 'issues_fixed', 'processing_time',
                        'error_message', 'profile_used', 'user',
                        'created_at', 'updated_at'
                    ]
                    
                    writer = csv.DictWriter(f, fieldnames=fieldnames)
                    writer.writeheader()
                    
                    # 데이터 작성
                    for history in histories:
                        row = history.to_dict()
                        # metadata 제외 (CSV에는 복잡한 구조 저장 어려움)
                        row.pop('metadata', None)
                        writer.writerow(row)
            
            self.logger.info(f"CSV 내보내기 완료: {output_path}")
            return True
            
        except Exception as e:
            self.logger.error(f"CSV 내보내기 실패: {e}")
            return False
    
    def _export_to_json(self, histories: List[ProcessHistory], output_path: Path) -> bool:
        """JSON 형식으로 내보내기"""
        try:
            with open(output_path, 'w', encoding='utf-8') as f:
                data = [h.to_dict() for h in histories]
                json.dump(data, f, ensure_ascii=False, indent=2, default=str)
            
            self.logger.info(f"JSON 내보내기 완료: {output_path}")
            return True
            
        except Exception as e:
            self.logger.error(f"JSON 내보내기 실패: {e}")
            return False
    
    def _export_to_excel(self, histories: List[ProcessHistory], output_path: Path) -> bool:
        """Excel 형식으로 내보내기"""
        try:
            import openpyxl
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill
            
            wb = Workbook()
            ws = wb.active
            ws.title = "처리 이력"
            
            # 헤더 스타일
            header_font = Font(bold=True)
            header_fill = PatternFill(start_color="D3D3D3", end_color="D3D3D3", fill_type="solid")
            
            # 헤더 작성
            headers = [
                'ID', '파일명', '파일 경로', '파일 크기', '처리 유형',
                '상태', '품질 점수', '발견 이슈', '수정 이슈',
                '처리 시간(초)', '오류 메시지', '프로파일', '사용자',
                '생성일시', '수정일시'
            ]
            
            for col, header in enumerate(headers, 1):
                cell = ws.cell(row=1, column=col, value=header)
                cell.font = header_font
                cell.fill = header_fill
            
            # 데이터 작성
            for row_idx, history in enumerate(histories, 2):
                data = history.to_dict()
                ws.cell(row=row_idx, column=1, value=data.get('id'))
                ws.cell(row=row_idx, column=2, value=data.get('file_name'))
                ws.cell(row=row_idx, column=3, value=data.get('file_path'))
                ws.cell(row=row_idx, column=4, value=data.get('file_size'))
                ws.cell(row=row_idx, column=5, value=data.get('process_type'))
                ws.cell(row=row_idx, column=6, value=data.get('status'))
                ws.cell(row=row_idx, column=7, value=data.get('quality_score'))
                ws.cell(row=row_idx, column=8, value=data.get('issues_found'))
                ws.cell(row=row_idx, column=9, value=data.get('issues_fixed'))
                ws.cell(row=row_idx, column=10, value=data.get('processing_time'))
                ws.cell(row=row_idx, column=11, value=data.get('error_message'))
                ws.cell(row=row_idx, column=12, value=data.get('profile_used'))
                ws.cell(row=row_idx, column=13, value=data.get('user'))
                ws.cell(row=row_idx, column=14, value=data.get('created_at'))
                ws.cell(row=row_idx, column=15, value=data.get('updated_at'))
            
            # 열 너비 자동 조정
            for column in ws.columns:
                max_length = 0
                column_letter = column[0].column_letter
                for cell in column:
                    try:
                        if len(str(cell.value)) > max_length:
                            max_length = len(str(cell.value))
                    except:
                        pass
                adjusted_width = min(max_length + 2, 50)
                ws.column_dimensions[column_letter].width = adjusted_width
            
            wb.save(output_path)
            self.logger.info(f"Excel 내보내기 완료: {output_path}")
            return True
            
        except ImportError:
            self.logger.error("openpyxl 라이브러리가 설치되지 않았습니다")
            return False
        except Exception as e:
            self.logger.error(f"Excel 내보내기 실패: {e}")
            return False
    
    def cleanup_old_history(self, days: int = 90) -> int:
        """
        오래된 이력 정리
        
        Args:
            days: 보관 일수
            
        Returns:
            int: 삭제된 레코드 수
        """
        cutoff_date = datetime.now() - timedelta(days=days)
        
        with self.db.get_connection() as conn:
            cursor = conn.cursor()
            
            # 오래된 이력 삭제
            cursor.execute("""
                DELETE FROM process_history 
                WHERE created_at < ?
            """, (cutoff_date.isoformat(),))
            
            deleted_history = cursor.rowcount
            
            # 오래된 통계도 정리
            cursor.execute("""
                DELETE FROM daily_statistics 
                WHERE date < ?
            """, (cutoff_date.date(),))
            
            deleted_stats = cursor.rowcount
            
            conn.commit()
        
        self.logger.info(
            f"오래된 데이터 정리 완료: "
            f"이력 {deleted_history}개, 통계 {deleted_stats}개 삭제 "
            f"(기준: {days}일)"
        )
        
        return deleted_history
    
    def archive_old_history(self, days: int = 365, archive_path: Optional[Path] = None) -> bool:
        """
        오래된 이력 아카이브
        
        Args:
            days: 아카이브 기준 일수
            archive_path: 아카이브 파일 경로
            
        Returns:
            bool: 성공 여부
        """
        cutoff_date = datetime.now() - timedelta(days=days)
        
        # 아카이브할 이력 조회
        histories = self.search.search_history(end_date=cutoff_date)
        
        if not histories:
            self.logger.info("아카이브할 이력이 없습니다")
            return True
        
        # 아카이브 경로 설정
        if archive_path is None:
            archive_dir = Path("data/archive")
            archive_dir.mkdir(parents=True, exist_ok=True)
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            archive_path = archive_dir / f"history_archive_{timestamp}.json"
        
        # JSON으로 저장
        if self._export_to_json(histories, archive_path):
            # 아카이브된 이력 삭제
            deleted = self.cleanup_old_history(days)
            self.logger.info(
                f"아카이브 완료: {len(histories)}개 이력 저장, "
                f"{deleted}개 삭제"
            )
            return True
        
        return False
    
    def vacuum_database(self) -> bool:
        """
        데이터베이스 최적화 (VACUUM)
        
        Returns:
            bool: 성공 여부
        """
        try:
            self.db.execute("VACUUM")
            self.logger.info("데이터베이스 최적화 완료")
            return True
        except Exception as e:
            self.logger.error(f"데이터베이스 최적화 실패: {e}")
            return False